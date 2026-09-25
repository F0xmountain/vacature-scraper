"""Export: nieuwe vacatures naar een xlsx per run, plus een cumulatieve csv.

De xlsx heeft een samenvatting plus de gevonden vaardigheden; de volledige
tekst blijft eruit, want die lees je niet in een cel. De csv houdt bewust de
oude negen kolommen, omdat hij wordt aangevuld en een kop van negen al vast
ligt.

Daarnaast laatste_run.json: dezelfde vacatures, maar met de volledige
omschrijving erbij. Zonder dat bestand werd die tekst na afloop weggegooid en
moest de webinterface alles opnieuw ophalen om hem te tonen: een halfuur
wachten op iets wat om 08:00 al binnen was.
"""
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

import vaardigheden


COLS = [
    "eerste_keer_gezien", "bedrijf", "functie", "locatie",
    "geplaatst", "salaris", "bron", "flags",
    "datasignalen", "vaardigheden", "samenvatting", "url",
]
HEADERS = [
    "Gezien", "Bedrijf", "Functie", "Locatie",
    "Geplaatst", "Salaris", "Bron", "Flags",
    "Datasignalen", "Vaardigheden", "Samenvatting", "URL",
]
BREEDTES = [11, 26, 44, 20, 12, 20, 20, 24, 12, 34, 70, 10]

# De csv houdt bewust de oude negen kolommen. Hij wordt aangevuld met
# mode="a" en de kop wordt alleen geschreven als het bestand nog niet
# bestaat; nieuwe kolommen toevoegen zou dus twaalf waarden onder een kop
# van negen plakken en het hele logboek stilletjes scheeftrekken.
CSV_COLS = [
    "eerste_keer_gezien", "bedrijf", "functie", "locatie",
    "geplaatst", "salaris", "bron", "flags", "url",
]


def _samenvatting(tekst, limiet=300):
    """Eerste stuk van de omschrijving, afgekapt op een woordgrens.

    De volledige tekst gaat bewust niet de xlsx in. Hij past wel (de langste
    is 15.861 tekens, ruim onder Excels celgrens van 32.767), maar 4.500
    tekens in een cel lees je alleen in de formulebalk. Drie compacte
    kolommen zeggen hier meer dan een muur proza: de eerste driehonderd
    tekens vertellen al waar de rol over gaat, en op vaardigheden en
    datasignalen kun je sorteren en filteren, wat een spreadsheet juist
    goed kan.
    """
    t = " ".join((tekst or "").split())
    if len(t) <= limiet:
        return t
    knip = t[:limiet]
    spatie = knip.rfind(" ")
    return (knip[:spatie] if spatie > 0 else knip).rstrip() + "..."



def export(nieuwe_jobs, cfg):
    outdir = Path(__file__).with_name(cfg.get("output_map") or "output")
    outdir.mkdir(exist_ok=True)

    df = pd.DataFrame(nieuwe_jobs)
    # Afgeleid uit de omschrijving, die zelf niet in de xlsx belandt.
    df["samenvatting"] = [_samenvatting(j.get("beschrijving")) for j in nieuwe_jobs]
    gevonden = [vaardigheden.uit_tekst(j.get("beschrijving")) for j in nieuwe_jobs]
    df["vaardigheden"] = [", ".join(v) for v in gevonden]
    df["datasignalen"] = [vaardigheden.kernsignalen(v) for v in gevonden]
    for c in COLS:
        if c not in df.columns:
            df[c] = ""
    df = df[COLS].fillna("").astype(str)
    df = df.replace({"NaT": "", "None": "", "nan": ""})
    # Terug naar een getal: astype(str) hierboven maakt er tekst van, en dan
    # sorteert Excel deze kolom alfabetisch in plaats van op grootte.
    df["datasignalen"] = pd.to_numeric(df["datasignalen"], errors="coerce").fillna(0).astype(int)

    df = df.sort_values("geplaatst", ascending=False)

    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    xlsx = outdir / f"nieuw_{stamp}.xlsx"

    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, header=HEADERS, sheet_name="Nieuw")
        ws = xw.sheets["Nieuw"]
        for i, w in enumerate(BREEDTES, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        url_col = COLS.index("url") + 1
        for rij in range(2, len(df) + 2):
            cel = ws.cell(row=rij, column=url_col)
            link = str(cel.value or "")
            if link.startswith("http"):
                cel.hyperlink = link
                cel.value = "open"
                cel.style = "Hyperlink"
        ws.freeze_panes = "A2"

    master = outdir / "alle_vacatures.csv"
    df[CSV_COLS].to_csv(
        master, mode="a", header=not master.exists(),
        index=False, encoding="utf-8-sig",
    )
    return xlsx


# Velden die de webinterface nodig heeft. Bewust een vaste lijst: zo lekken
# interne vlaggen (zoals _detail) niet mee en blijft het bestand voorspelbaar.
# Hoeveel runs het venster bewaart. Zes is drie dagen bij twee runs per dag.
BEWAAR_RUNS = 6

BEWAAR_VELDEN = (
    "bron", "bedrijf", "functie", "locatie", "geplaatst",
    "salaris", "flags", "url", "beschrijving", "eerste_keer_gezien",
)


def bewaar_run(nieuwe_jobs, cfg, xlsx=None):
    """Voeg deze run toe aan output/laatste_run.json en kap het venster af.

    Bewust een venster en niet een enkele run. Met twee runs per dag zou de
    avondrun de ochtendoogst overschrijven voordat je die gelezen had, en ook bij
    een run per dag verdween een batch zodra je de app een dag niet opende. De
    xlsx-bestanden blijven daarnaast gewoon staan als archief.

    Een lege run wordt ook toegevoegd: dan blijft zichtbaar dat hij gedraaid heeft
    en niets vond, terwijl de eerdere runs in het venster gewoon leesbaar blijven.
    """
    outdir = Path(__file__).with_name(cfg.get("output_map") or "output")
    outdir.mkdir(exist_ok=True)
    pad = outdir / "laatste_run.json"

    bestaand = []
    if pad.exists():
        try:
            bestaand = (json.loads(pad.read_text(encoding="utf-8")) or {}).get("runs") or []
        except (ValueError, OSError):
            bestaand = []

    bestaand.append({
        "moment": datetime.now().isoformat(timespec="seconds"),
        "bestand": str(xlsx) if xlsx else None,
        "aantal": len(nieuwe_jobs),
        "jobs": [{v: (j.get(v) or "") for v in BEWAAR_VELDEN} for j in nieuwe_jobs],
    })
    pad.write_text(
        json.dumps({"runs": bestaand[-BEWAAR_RUNS:]}, ensure_ascii=False, indent=1),
        encoding="utf-8",
    )
    return pad


def lees_laatste_run(cfg):
    """Het venster teruglezen als een platte lijst, of None als er nog niets is.

    De runs worden samengevoegd, nieuwste eerst, en ontdubbeld op URL; met een
    kort terugkijkvenster overlappen opeenvolgende runs namelijk. moment is dat
    van de nieuwste run.
    """
    pad = (
        Path(__file__).with_name(cfg.get("output_map") or "output")
        / "laatste_run.json"
    )
    if not pad.exists():
        return None
    try:
        runs = (json.loads(pad.read_text(encoding="utf-8")) or {}).get("runs") or []
    except (ValueError, OSError):
        return None
    if not runs:
        return None

    jobs, gezien = [], set()
    for run in reversed(runs):
        for j in run.get("jobs") or []:
            sleutel = j.get("url") or repr(sorted(j.items()))
            if sleutel in gezien:
                continue
            gezien.add(sleutel)
            jobs.append(j)

    return {
        "moment": runs[-1].get("moment"),
        "aantal": len(jobs),
        "runs": len(runs),
        "jobs": jobs,
    }
